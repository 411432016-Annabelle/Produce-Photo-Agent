"""The queue worker end to end, with a fake Discord message and a fake claude command."""

import asyncio

from openpyxl import load_workbook

from agent import excel_log
from agent.queue_worker import Job, MessageTracker, Worker
from tests.helpers import BANANA, CAT, jpeg_bytes, make_fake_claude


class FakeAttachment:
    def __init__(self, data: bytes):
        self.data = data

    async def read(self):
        return self.data


class FakeMessage:
    def __init__(self, fail_reply=False):
        self.replies, self.removed, self.fail_reply = [], [], fail_reply

    async def reply(self, text, **kwargs):
        if self.fail_reply:
            raise RuntimeError("message was deleted")
        self.replies.append((text, kwargs))

    async def remove_reaction(self, emoji, user):
        self.removed.append(emoji)


def run_jobs(tmp_path, answer, photos, caption="bought yesterday", message=None):
    data_dir = tmp_path / "data"
    claude = make_fake_claude(tmp_path, answer, tmp_path / "record.json")
    message = message or FakeMessage()

    async def main():
        worker = Worker(data_dir, data_dir / "photos", "haiku", claude_bin=claude)
        tracker = MessageTracker(remaining=len(photos))
        for index, data in enumerate(photos, start=1):
            await worker.queue.put(Job(message, FakeAttachment(data), index, len(photos), caption,
                                       "=Anna", "123456789012345678", "https://discord.com/x", tracker))
        task = asyncio.create_task(worker.run_forever())
        await worker.queue.join()
        task.cancel()
        return worker

    worker = asyncio.run(main())
    rows = list(load_workbook(worker.excel.path)["Log"].iter_rows(min_row=2, values_only=True))
    return message, rows, sorted(p.name for p in (data_dir / "photos").glob("*.jpg"))


def test_produce_photo(tmp_path):
    message, rows, files = run_jobs(tmp_path, BANANA, [jpeg_bytes()])
    text, kwargs = message.replies[0]
    assert text.startswith("🍎 Banana (fruit)") and text.endswith("Today: 1 checked (1 fruits, 0 vegetables)")
    assert kwargs["mention_author"] is False and kwargs["allowed_mentions"].everyone is False
    assert message.removed == ["⏳"]
    row = rows[0]
    assert row[1:8] == ("=Anna", "123456789012345678", "fruit", "Banana", 72, "Yes", 2)
    assert row[11] == "bought yesterday" and row[12].endswith("Banana.jpg") and row[14:] == (1, 1, 0)
    assert len(files) == 1 and files[0].endswith(" Banana.jpg")


def test_not_produce_is_logged_but_not_counted(tmp_path):
    message, rows, files = run_jobs(tmp_path, CAT, [jpeg_bytes()])
    assert message.replies[0][0].startswith("This doesn't look like a fruit or vegetable — I see a cat.")
    assert rows[0][4] == "Not produce" and rows[0][9] == "a cat" and rows[0][14:] == (0, 0, 0)
    assert files[0].endswith(" Not produce.jpg")


def test_invalid_answer_gives_error_row(tmp_path):
    message, rows, files = run_jobs(tmp_path, {**BANANA, "ready_to_eat": "ripe"}, [jpeg_bytes()])
    assert message.replies[0][0].startswith("Recognition failed, please send the photo again")
    assert rows[0][4] == "Error" and rows[0][9] == "ready_to_eat is not yes/not_yet/past_best"
    assert rows[0][14:] == (0, 0, 0) and files[0].endswith(" Error.jpg")


def test_broken_photo_gives_error_without_file(tmp_path):
    message, rows, files = run_jobs(tmp_path, BANANA, [b"not an image"])
    assert message.replies[0][0].startswith("Recognition failed")
    assert rows[0][4] == "Error" and rows[0][12] is None and files == []


def test_three_photos_three_replies_one_reaction_removal(tmp_path):
    message, rows, files = run_jobs(tmp_path, BANANA, [jpeg_bytes()] * 3)
    assert [t.splitlines()[0] for t, _ in message.replies] == ["Photo 1 of 3", "Photo 2 of 3", "Photo 3 of 3"]
    assert [r[14] for r in rows] == [1, 2, 3]
    assert len(files) == 3
    assert message.removed == ["⏳"]


def test_failed_reply_never_crashes(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_log, "RETRY_SECONDS", 0)
    message, rows, _ = run_jobs(tmp_path, BANANA, [jpeg_bytes()], message=FakeMessage(fail_reply=True))
    assert len(rows) == 1 and message.removed == ["⏳"]
