import asyncio
import zipfile
from datetime import date

from openpyxl import load_workbook

from agent import excel_log
from agent.excel_log import COLUMNS, ExcelLog

TODAY = date(2026, 10, 26)


def row(time="2026-10-26 10:00:00", name="Banana", produce_type="fruit", display="Anna", note=""):
    return [time, display, "123456789012345678", produce_type, name, 72, "Yes", 2, "Tip",
            "Desc", 88, note, "/p.jpg", "https://discord.com/channels/1/2/3", 1, 1, 0]


def sheet_xml(path):
    with zipfile.ZipFile(path) as archive:
        return archive.read("xl/worksheets/sheet1.xml").decode()


def test_creates_header_and_appends(tmp_path):
    log = ExcelLog(tmp_path)
    asyncio.run(log.append(row(name="Banana")))
    asyncio.run(log.append(row(name="Carrot", produce_type="vegetable")))
    sheet = load_workbook(log.path)["Log"]
    values = list(sheet.iter_rows(values_only=True))
    assert values[0] == COLUMNS
    assert [v[4] for v in values[1:]] == ["Banana", "Carrot"]
    assert values[1][2] == "123456789012345678"  # user ID kept as text, all digits


def test_no_formulas_are_ever_written(tmp_path):
    log = ExcelLog(tmp_path)
    asyncio.run(log.append(row(display='=HYPERLINK("http://evil.example","x")', note="@SUM(1)")))
    asyncio.run(log.append(row(display="+1", note="-2", name="\tTab")))
    assert "<f>" not in sheet_xml(log.path)
    sheet = load_workbook(log.path)["Log"]
    assert sheet["B2"].value.startswith("=HYPERLINK")
    assert sheet["B2"].data_type == "s" and sheet["B2"].quotePrefix
    assert sheet["L2"].quotePrefix and sheet["B3"].quotePrefix and sheet["L3"].quotePrefix


def test_busy_file_goes_to_backup_then_recovers(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_log, "RETRY_SECONDS", 0)
    log = ExcelLog(tmp_path)
    real_write = log._write_rows

    def locked(rows):
        raise PermissionError("file is locked")

    monkeypatch.setattr(log, "_write_rows", locked)
    asyncio.run(log.append(row(name="Apple")))
    assert log.pending_path.exists() and not log.path.exists()

    monkeypatch.setattr(log, "_write_rows", real_write)
    asyncio.run(log.append(row(name="Pear")))
    names = [r[4] for r in load_workbook(log.path)["Log"].iter_rows(min_row=2, values_only=True)]
    assert names == ["Apple", "Pear"]
    assert not log.pending_path.exists()


def test_today_counts_survive_restart(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_log, "RETRY_SECONDS", 0)
    log = ExcelLog(tmp_path)
    asyncio.run(log.append(row(name="Banana")))
    asyncio.run(log.append(row(name="Carrot", produce_type="vegetable")))
    asyncio.run(log.append(row(name="Not produce", produce_type="")))
    asyncio.run(log.append(row(name="Error", produce_type="")))
    asyncio.run(log.append(row(time="2026-10-25 23:59:59", name="Old")))
    log.pending_path.write_text('["2026-10-26 11:00:00","","","fruit","Kiwi"]\n')

    counts = ExcelLog(tmp_path).count_today(TODAY)
    assert (counts.total, counts.fruits, counts.vegetables) == (3, 2, 1)


def test_counts_reset_on_new_day():
    counts = excel_log.DailyCounts(TODAY, 5, 3, 2)
    counts.roll_to(TODAY)
    assert counts.total == 5
    counts.roll_to(date(2026, 10, 27))
    assert (counts.total, counts.fruits, counts.vegetables) == (0, 0, 0)
