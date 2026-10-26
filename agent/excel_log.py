"""The append-only Excel log, today's counts, and the busy-file backup (architecture section 8)."""

import asyncio
import json
import os
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .log import log

SHEET_NAME = "Log"
COLUMNS = (
    "Time", "Display name", "User ID", "Type", "Item name", "Freshness score",
    "Ready to eat", "Days left", "Storage tip", "Description", "Confidence", "Note",
    "Photo file", "Message link", "Today total", "Today fruits", "Today vegetables",
)
NOT_PRODUCE = "Not produce"
ERROR = "Error"
SAVE_ATTEMPTS = 3
RETRY_SECONDS = 2
FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r", "\n")


@dataclass
class DailyCounts:
    day: date
    total: int = 0
    fruits: int = 0
    vegetables: int = 0

    def roll_to(self, today: date) -> None:
        if today != self.day:
            self.day, self.total, self.fruits, self.vegetables = today, 0, 0, 0

    def add(self, produce_type: str) -> None:
        self.total += 1
        if produce_type == "fruit":
            self.fruits += 1
        else:
            self.vegetables += 1


class ExcelLog:
    def __init__(self, data_dir: Path):
        self.path = data_dir / "produce_log.xlsx"
        self.pending_path = data_dir / "pending_rows.jsonl"
        self.temp_path = data_dir / "produce_log.saving.xlsx"
        data_dir.mkdir(parents=True, exist_ok=True)

    def count_today(self, today: date) -> DailyCounts:
        """Rebuilds today's counts from the sheet and any pending rows, so a restart keeps them."""
        counts = DailyCounts(today)
        rows = list(self._read_pending())
        if self.path.exists():
            try:
                sheet = load_workbook(self.path, read_only=True)[SHEET_NAME]
                rows += [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)]
            except Exception as error:
                log(f"⚠️  Could not read {self.path.name} to count today's rows ({error}). Starting from 0.")
        prefix = today.isoformat()
        for row in rows:
            time, produce_type, name = row[0], row[3], row[4]
            if (
                isinstance(time, str) and time.startswith(prefix)
                and name not in (NOT_PRODUCE, ERROR)
                and produce_type in ("fruit", "vegetable")
            ):
                counts.add(produce_type)
        return counts

    async def append(self, row: list) -> None:
        """Adds one row, plus any rows waiting in the backup file. Never raises."""
        if (self.path.parent / f"~${self.path.name}").exists():
            log(f"⚠️  {self.path.name} looks open in Excel. Close it without saving, or new rows may be overwritten.")
        rows = list(self._read_pending()) + [row]
        for attempt in range(1, SAVE_ATTEMPTS + 1):
            try:
                await asyncio.to_thread(self._write_rows, rows)
                self.pending_path.unlink(missing_ok=True)
                return
            except Exception as error:
                log(f"⚠️  Could not save {self.path.name} (is it open in Excel?) "
                    f"[attempt {attempt}/{SAVE_ATTEMPTS}: {error}]. Retrying…")
                if attempt < SAVE_ATTEMPTS:
                    await asyncio.sleep(RETRY_SECONDS)
        try:
            with open(self.pending_path, "a", encoding="utf-8") as file:
                file.write(json.dumps(row, ensure_ascii=False) + "\n")
            log(f"⚠️  Row kept in {self.pending_path.name} — it will be added to Excel "
                "automatically once the file can be saved.")
        except Exception as error:
            log(f"❌ Could not write the backup file either ({error}). Row: {row}")

    def _read_pending(self):
        if not self.pending_path.exists():
            return
        try:
            lines = self.pending_path.read_text(encoding="utf-8").splitlines()
        except Exception as error:
            log(f"⚠️  Could not read {self.pending_path.name} ({error}).")
            return
        for line in lines:
            if line.strip():
                try:
                    yield json.loads(line)
                except json.JSONDecodeError:
                    log(f"⚠️  Skipping an unreadable line in {self.pending_path.name}: {line[:80]}")

    def _write_rows(self, rows: list[list]) -> None:
        if self.path.exists():
            workbook = load_workbook(self.path)
            sheet = workbook[SHEET_NAME]
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = SHEET_NAME
            for column, header in enumerate(COLUMNS, start=1):
                sheet.cell(row=1, column=column, value=header)
        for row in rows:
            row_number = sheet.max_row + 1
            for column, value in enumerate(row, start=1):
                cell = sheet.cell(row=row_number, column=column)
                if isinstance(value, str):
                    write_text(cell, value)
                else:
                    cell.value = value
        # Save to a side file first, then swap it in, so a half-written file
        # never replaces the real log.
        workbook.save(self.temp_path)
        os.replace(self.temp_path, self.path)


def write_text(cell, text: str) -> None:
    cell.value = text

    # ==== GUARD 4: Excel formula protection START ====
    # Text that Excel could read as a formula is stored as plain text: the cell is
    # forced to the text type and given Excel's hidden apostrophe ("quote prefix").
    if text.startswith(FORMULA_STARTS):
        cell.data_type = "s"
        cell.quotePrefix = True
    # ==== GUARD 4 END ====
